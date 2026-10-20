import io
import unittest

from docx import Document
from streamlit.testing.v1 import AppTest

import streamlit_app


def render_tor_test_app():
    import streamlit_app as app_module
    app_module.render_tor(("RAB", app_module.st.session_state["test_analysis"]))


class StreamlitAppTests(unittest.TestCase):
    def test_dashboard_opens_without_login_or_signup(self):
        app = AppTest.from_file(streamlit_app.__file__).run(timeout=20)

        self.assertEqual(len(app.exception), 0)
        self.assertEqual([title.value for title in app.title], ["Dashboard RAB & Generator TOR"])
        self.assertEqual(len(app.get("file_uploader")), 1)
        self.assertFalse(any("Masuk" in item.value or "Daftar" in item.value for item in app.markdown))

    def test_to_number_handles_indonesian_and_plain_amounts(self):
        self.assertEqual(streamlit_app.to_number("Rp 1.234.567"), 1234567)
        self.assertEqual(streamlit_app.to_number("1,234,567"), 1234567)
        self.assertEqual(streamlit_app.to_number("1234,5"), 1234.5)
        self.assertIsNone(streamlit_app.to_number("—"))

    def test_compact_rupiah_uses_short_indonesian_units(self):
        self.assertEqual(streamlit_app.format_compact_rupiah(17_100_000_000), "17,1 M")
        self.assertEqual(streamlit_app.format_compact_rupiah(1_260_000), "1,3 jt")

    def test_revision_compares_remaining_to_increase(self):
        rows = [
            ["", "Semula", None, None, "Menjadi", None, None, None, None],
            [None, "Kode Anggaran", "Uraian", "Jumlah", "Kode Anggaran", "Uraian", "Jumlah", "Selisih", "Sisa Anggaran"],
            [None, None, "Perjalanan", 100, None, "Perjalanan", 120, 20, 15],
            [None, None, "Peralatan", 100, None, "Peralatan", 110, 10, 25],
        ]
        revision = streamlit_app.find_revision(rows)
        self.assertEqual(revision["increase_total"], 30)
        self.assertEqual(revision["insufficient_count"], 1)
        self.assertEqual(revision["sufficient_count"], 1)
        self.assertEqual(revision["shortfall"], 5)

    def test_dashboard_analysis_prepares_rows_and_categories(self):
        rows = [
            ["Kode Anggaran", "Uraian", "Pagu", "Sisa Anggaran"],
            ["051", "Monitoring", 100000, 20000],
            ["052", "Pengawasan", 80000, 10000],
            ["", "Baris tidak bernilai", None, None],
        ]
        analysis = streamlit_app.analyze_rows(rows, "RAB")
        self.assertTrue(analysis["header_found"])
        self.assertEqual(len(analysis["rows"]), 2)
        self.assertEqual(len(analysis["categories"]), 2)
        self.assertEqual(analysis["category_total"], 180000)

    def test_sp2d_pdf_text_parser_reads_category_and_wrapped_detail_rows(self):
        pages = ["""
LAPORAN REALISASI SP2D TA 2026
                       Uraian                    Pagu Revisi Lock Pagu
     100 Kegiatan pertama                     1,000,000 0 500,000 0 500,000 50.00 % 500,000
       101 Kegiatan kedua                      2,000,000 0 1,000,000 0 1,000,000 50.00 % 1,000,000
              000001. Rincian kegiatan yang
                                          100,000 0 50,000 0 50,000 50.00 % 50,000
                 keterangannya terbungkus
"""]

        rows = streamlit_app.parse_sp2d_text(pages)

        self.assertEqual(rows[0], ["Kode Anggaran", "Uraian", "Pagu", "Sisa Anggaran"])
        self.assertEqual(rows[1], ["100", "Kegiatan pertama", 1000000, 500000])
        self.assertEqual(rows[2], ["101", "Kegiatan kedua", 2000000, 1000000])
        self.assertEqual(rows[3], ["000001", "Rincian kegiatan yang keterangannya terbungkus", 100000, 50000])
        self.assertIsNone(streamlit_app.parse_sp2d_text(["PDF laporan lain"]))

    def test_budget_pdf_parsers_combine_ro_budget_and_financial_values(self):
        workplan = streamlit_app.parse_rincian_kertas_text(["""
RINCIAN KERTAS KERJA SATKER T.A. 2026
ALOKASI Rp. 1,500,000
7437.BAH.078 Layanan Dukungan RO 1.0 layanan 1,500,000
- Rincian terblokir 100,000 100,000 *
"""])
        financial = streamlit_app.parse_fa_text(["""
LAPORAN REALISASI SP2D TA 2026
JUMLAH SELURUHNYA 1,000,000 0 400,000 0 400,000 40.00 % 600,000
BAH.078 Layanan Dukungan RO 1,000,000 0 400,000 0 400,000 40.00 % 600,000
"""])

        rows = streamlit_app.combine_budget_documents([
            ("Rincian Kertas Kerja", workplan),
            ("Laporan FA Detail", financial),
        ])
        analysis = streamlit_app.analyze_rows(rows, "Perbandingan Anggaran RO")

        self.assertEqual(analysis["budget_comparison"]["pagu_penuh"], 1500000)
        self.assertEqual(analysis["budget_comparison"]["pagu_efektif"], 1000000)
        self.assertEqual(analysis["budget_comparison"]["realisasi"], 400000)
        self.assertEqual(analysis["budget_comparison"]["sisa"], 600000)
        self.assertEqual(analysis["budget_comparison"]["pagu_terblokir"], 100000)
        self.assertEqual(analysis["budget_comparison"]["matched_count"], 1)
        self.assertEqual(analysis["budget_comparison"]["unmatched_count"], 0)
        ro = analysis["rows"][0]
        self.assertEqual(ro["Kode"], "7437.BAH.078")
        self.assertEqual(ro["Pagu Terblokir"], 100000)
        self.assertEqual(ro["Status Pasangan"], "Cocok berdasarkan kode RO")
        sections = streamlit_app.build_tor_sections(
            ro["Uraian"], "Surabaya", "2026", "", analysis, selected_row=ro,
        )
        financing = next(body for title, body in sections if title == "VII. Pembiayaan")
        self.assertIn("Kode RO: 7437.BAH.078", financing)
        self.assertIn("Pagu penuh: Rp 1.500.000", financing)
        self.assertIn("Pagu terblokir: Rp 100.000", financing)

    def test_budget_parser_and_analysis_preserve_full_hierarchy(self):
        workplan = streamlit_app.parse_rincian_kertas_text(["""
RINCIAN KERTAS KERJA SATKER T.A. 2026
ALOKASI Rp. 1,500,000
059.09.GL Dukungan layanan 1.0 program 1,500,000
7437 Kegiatan dukungan 1.0 kegiatan 1,500,000
7437.BAH Layanan dukungan 1.0 layanan 1,500,000
7437.BAH.078 Layanan pengujian 1.0 layanan 1,500,000
101 Komponen utama 1.0 paket 1,500,000
A Perjalanan dinas 1.0 paket 1,500,000
521115 Belanja perjalanan 1.0 paket 1,500,000
- Tiket perjalanan 1.0 paket 1,500,000 *
"""])

        analysis = streamlit_app.analyze_rows(workplan, "Rincian Kertas Kerja")
        rows_by_level = {row["Tingkat"]: row for row in analysis["rows"]}

        self.assertEqual(
            set(rows_by_level),
            {"Program", "Kegiatan", "KRO", "RO", "Komponen", "Subkomponen", "Detil belanja"},
        )
        self.assertEqual(rows_by_level["Program"]["Kode"], "059.09.GL")
        self.assertEqual(rows_by_level["Kegiatan"]["Kode"], "7437")
        self.assertEqual(rows_by_level["Komponen"]["Kode"], "101")
        self.assertEqual(rows_by_level["Subkomponen"]["Kode"], "A")
        self.assertEqual(rows_by_level["Detil belanja"]["Kode"], "521115")
        self.assertEqual(rows_by_level["Detil belanja"]["Pagu Terblokir"], 1500000)
        self.assertEqual(rows_by_level["Kegiatan"]["Induk"], "059.09.GL")
        self.assertEqual(rows_by_level["Detil belanja"]["Induk"], "A")
        self.assertEqual(analysis["hierarchy_counts"]["DETIL"], 1)
        self.assertIn("DETIL", analysis["category_levels"])
        self.assertNotIn("AKUN", analysis["category_levels"])
        all_categories = analysis["category_levels"]["ALL"]
        self.assertEqual(
            len(all_categories),
            sum(len(categories) for level, categories in analysis["category_levels"].items() if level != "ALL"),
        )
        self.assertEqual(
            {category["Tingkat"] for category in all_categories},
            {"Program", "Kegiatan", "KRO", "RO", "Komponen", "Subkomponen", "Detil belanja"},
        )

    def test_tor_uses_template_and_edits_only_requested_sections(self):
        workplan = streamlit_app.parse_rincian_kertas_text(["""
RINCIAN KERTAS KERJA SATKER T.A. 2026
ALOKASI Rp. 1,500,000
059.09.GL Program utama 1.0 program 1,500,000
7437 Kegiatan utama 1.0 kegiatan 1,500,000
7437.BAH KRO utama 1.0 layanan 1,500,000
7437.BAH.078 RO terpilih 1.0 layanan 1,500,000
101 Komponen utama 1.0 paket 1,500,000
A Subkomponen utama 1.0 paket 1,500,000
521115 Detil perjalanan 1.0 paket 1,500,000
- Rincian terblokir 100,000 100,000 *
"""])
        financial = streamlit_app.parse_fa_text(["""
LAPORAN REALISASI SP2D TA 2026
JUMLAH SELURUHNYA 1,000,000 0 400,000 0 400,000 40.00 % 600,000
BAH.078 RO terpilih 1,000,000 0 400,000 0 400,000 40.00 % 600,000
"""])
        analysis = streamlit_app.analyze_rows(
            streamlit_app.combine_budget_documents([
                ("Rincian Kertas Kerja", workplan),
                ("Laporan FA Detail", financial),
            ]),
            "RAB",
        )
        selected = next(row for row in analysis["rows"] if row["Tingkat"] == "RO")
        hierarchy, descendants = streamlit_app.get_tor_hierarchy(analysis, selected)
        template = Document(streamlit_app.find_local_tor_template())
        beneficiaries = streamlit_app.docx_table_frame(template.tables[4])
        beneficiaries.iloc[0, 0] = "Penerima manfaat RO terpilih"
        schedule = streamlit_app.docx_table_frame(
            template.tables[5],
            first_data_row=2,
            column_names=["No", "Rincian Aktivitas"] + ["Bulan {}".format(month) for month in range(1, 13)],
        )
        schedule.iloc[0, 1] = "Pelaksanaan RO terpilih"
        schedule.iloc[0, 2] = "X"
        output_kpi = streamlit_app.docx_table_frame(template.tables[6])
        output_kpi.iloc[0, 0] = "KPI output RO terpilih"
        outcome_kpi = streamlit_app.docx_table_frame(template.tables[7])
        outcome_kpi.iloc[0, 0] = "KPI outcome RO terpilih"
        adjustments = {
            "beneficiaries": streamlit_app.docx_frame_rows(beneficiaries),
            "strategy": "Strategi RO terpilih\n- Detil 521115 — Detil perjalanan",
            "schedule": streamlit_app.docx_frame_rows(schedule),
            "kpi_output": streamlit_app.docx_frame_rows(output_kpi),
            "kpi_outcome": streamlit_app.docx_frame_rows(outcome_kpi),
            "cost_benefit": [[
                "Biaya RO 7437.BAH.078: Rp 1.000.000",
                "Manfaat RO terpilih",
            ]],
            "risks": [[
                "1", "Risiko terpilih", "Penyebab terpilih", "Dampak terpilih",
                "Mitigasi terpilih", "PIC terpilih",
            ]],
        }

        content = streamlit_app.make_tor_docx(
            selected["Kode"],
            selected,
            hierarchy,
            adjustments,
        )
        document = Document(io.BytesIO(content))
        template_original = Document(streamlit_app.find_local_tor_template())

        self.assertEqual(len(document.tables), len(template_original.tables))
        self.assertEqual(document.sections[0].page_width, template_original.sections[0].page_width)
        self.assertEqual(document.sections[0].left_margin, template_original.sections[0].left_margin)
        self.assertEqual(document.tables[0].cell(0, 1).text, "059.09.GL")
        self.assertEqual(document.tables[0].cell(0, 2).text, "Program utama")
        self.assertEqual(document.tables[0].cell(1, 1).text, "7437")
        self.assertEqual(document.tables[0].cell(1, 2).text, "Kegiatan utama")
        self.assertEqual(document.tables[0].cell(2, 1).text, "7437.BAH")
        self.assertEqual(document.tables[0].cell(2, 2).text, "KRO utama")
        self.assertEqual(document.tables[0].cell(3, 1).text, "7437.BAH.078")
        self.assertEqual(document.tables[4].cell(1, 0).text, "Penerima manfaat RO terpilih")
        self.assertEqual(document.tables[5].cell(2, 1).text, "Pelaksanaan RO terpilih")
        self.assertEqual(document.tables[5].cell(2, 2).text, "X")
        self.assertEqual(document.tables[6].cell(1, 0).text, "KPI output RO terpilih")
        self.assertEqual(document.tables[7].cell(1, 0).text, "KPI outcome RO terpilih")
        self.assertEqual(len(document.tables[8].rows), len(template_original.tables[8].rows))
        self.assertEqual(document.tables[8].cell(1, 0).text, adjustments["cost_benefit"][0][0])
        self.assertEqual(document.tables[8].cell(1, 1).text, "Manfaat RO terpilih")
        self.assertEqual(
            [[cell.text for cell in row.cells] for row in document.tables[8].rows[2:]],
            [[cell.text for cell in row.cells] for row in template_original.tables[8].rows[2:]],
        )
        self.assertEqual(len(document.tables[9].rows), len(template_original.tables[9].rows))
        self.assertEqual(document.tables[9].cell(1, 1).text, "Risiko terpilih")
        self.assertEqual(
            [[cell.text for cell in row.cells] for row in document.tables[9].rows[2:]],
            [[cell.text for cell in row.cells] for row in template_original.tables[9].rows[2:]],
        )
        self.assertIn("Rp 1.000.000", next(
            paragraph.text for paragraph in document.paragraphs
            if paragraph.text.startswith("Biaya yang diperlukan untuk kegiatan ini sebesar")
        ))
        self.assertIn("Strategi RO terpilih", [paragraph.text for paragraph in document.paragraphs])
        self.assertTrue(descendants)
        self.assertIn(
            "Gambaran Umum",
            [paragraph.text for paragraph in document.paragraphs],
        )
        def static_paragraphs(doc):
            result = []
            in_strategy = False
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if text == "Metode Pelaksanaan":
                    in_strategy = True
                    result.append(paragraph.text)
                    continue
                if text == "Tahapan dan Waktu Pelaksanaan":
                    in_strategy = False
                if in_strategy or text.startswith("Biaya yang diperlukan untuk kegiatan ini sebesar"):
                    continue
                result.append(paragraph.text)
            return result

        self.assertEqual(static_paragraphs(document), static_paragraphs(template_original))
        for table_index in (1, 2, 3):
            self.assertEqual(
                [[cell.text for cell in row.cells] for row in document.tables[table_index].rows],
                [[cell.text for cell in row.cells] for row in template_original.tables[table_index].rows],
            )
        for table_index in (4, 5, 6, 7, 8, 9):
            self.assertEqual(
                [cell._tc.tcPr.xml for cell in document.tables[table_index].rows[0].cells],
                [cell._tc.tcPr.xml for cell in template_original.tables[table_index].rows[0].cells],
            )

    def test_tor_adjustment_controls_render_for_a_selected_ro(self):
        workplan = streamlit_app.parse_rincian_kertas_text(["""
RINCIAN KERTAS KERJA SATKER T.A. 2026
ALOKASI Rp. 1,500,000
059.09.GL Program utama 1.0 program 1,500,000
7437 Kegiatan utama 1.0 kegiatan 1,500,000
7437.BAH KRO utama 1.0 layanan 1,500,000
7437.BAH.078 RO terpilih 1.0 layanan 1,500,000
101 Komponen utama 1.0 paket 1,500,000
A Subkomponen utama 1.0 paket 1,500,000
521115 Detil perjalanan 1.0 paket 1,500,000
"""])
        financial = streamlit_app.parse_fa_text(["""
LAPORAN REALISASI SP2D TA 2026
JUMLAH SELURUHNYA 1,000,000 0 400,000 0 400,000 40.00 % 600,000
BAH.078 RO terpilih 1,000,000 0 400,000 0 400,000 40.00 % 600,000
"""])
        rows = streamlit_app.combine_budget_documents([
            ("Rincian Kertas Kerja", workplan),
            ("Laporan FA Detail", financial),
        ])
        analysis = streamlit_app.analyze_rows(rows, "RAB")
        app = AppTest.from_function(render_tor_test_app)
        app.session_state["test_analysis"] = analysis
        app.run(timeout=30)

        self.assertEqual(len(app.exception), 0)
        self.assertTrue(any(
            item.label == "Pilih kode Rincian Output (RO)"
            for item in app.selectbox
        ))
        next(button for button in app.button if button.label == "Buat draft TOR dari template").click()
        app.run(timeout=30)
        self.assertEqual(len(app.exception), 0)
        self.assertTrue(any(
            item.label == "Unduh TOR dalam Word"
            for item in app.get("download_button")
        ))


if __name__ == "__main__":
    unittest.main()
